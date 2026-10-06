"""Read-only, privacy-preserving inspection of the reviewed schedule layout."""
from datetime import datetime, time
from pathlib import Path
import re

from openpyxl import load_workbook

from app.parsers.curriculum_workbook import SUBJECT_SHEETS, parse_workbook

DAY_COLUMNS = {3: ("Monday", "RFBT"), 6: ("Tuesday", "AFAR"),
               9: ("Wednesday", None), 12: ("Thursday", "MAS"),
               15: ("Friday", "TAX"), 18: ("Saturday", "FAR"),
               21: ("Sunday", None)}


def inspect_schedule(path: Path) -> dict:
    parsed = parse_workbook(path)
    if parsed.errors:
        raise ValueError("Curriculum validation failed; schedule inspection is not approved.")
    topics = {(t["sheet"], t["row"]): t for t in parsed.topics}
    book = load_workbook(path, data_only=False, keep_links=False)
    try:
        if not {"SCHEDULE", "Calendar"}.issubset(book.sheetnames):
            raise ValueError("The reviewed SCHEDULE/Calendar sheets are required.")
        schedule, calendar = book["SCHEDULE"], book["Calendar"]
        blocks = [row for row in range(6, schedule.max_row + 1)
                  if all(isinstance(schedule.cell(row, column).value, datetime)
                         for column in DAY_COLUMNS)]
        if not blocks:
            raise ValueError("The reviewed seven-day schedule layout was not found.")
        direct = mapped = 0
        invalid, crossed, indirect = [], [], []
        exact_times = sum(schedule.cell(row, col).value.time() != time()
                          for row in blocks for col in DAY_COLUMNS)
        # Topic rows extend beyond three-row blocks in several weeks. Inspect all source
        # reference rows; target-date/countdown helper formulas cannot match this syntax.
        for start, stop in [(min(blocks), schedule.max_row + 1)]:
            for row in range(start, stop):
                for column, (_, expected) in DAY_COLUMNS.items():
                    cell = schedule.cell(row, column)
                    if cell.data_type != "f":
                        continue
                    match = re.fullmatch(r"=(?:'([^']+)'|([A-Z]+))!\$?C\$?(\d+)", cell.value)
                    if match:
                        name, topic_row = match.group(1) or match.group(2), int(match.group(3))
                        if name not in SUBJECT_SHEETS:
                            invalid.append(cell.coordinate)
                            continue
                        direct += 1
                        topic = topics.get((name, topic_row))
                        if topic is None:
                            invalid.append(cell.coordinate)
                            continue
                        mapped += 1
                        if expected and topic["subject"] != expected:
                            crossed.append((cell.coordinate, expected, topic["subject"], topic["code"]))
                    elif re.fullmatch(r"=\$?[A-Z]+\$?\d+", cell.value):
                        indirect.append(cell.coordinate)
        month_rows = [(row, calendar.cell(row, 2).value)
                      for row in range(1, calendar.max_row + 1)
                      if isinstance(calendar.cell(row, 2).value, datetime)
                      and calendar.cell(row, 2).value.day == 1]
        days = texts = unassigned = 0
        errors = []
        seen = set()
        for index, (start, month) in enumerate(month_rows):
            stop = month_rows[index + 1][0] if index + 1 < len(month_rows) else calendar.max_row + 1
            exact_times += int(month.time() != time())
            context = {}
            for row in range(start + 3, stop):
                # A new week must not inherit dates from blank columns in the preceding week.
                if any(isinstance(calendar.cell(row, c).value, (int, float))
                       and not isinstance(calendar.cell(row, c).value, bool) for c in range(2, 9)):
                    context = {}
                for column in range(2, 9):
                    cell = calendar.cell(row, column)
                    value = cell.value
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        try:
                            if value != int(value):
                                raise ValueError
                            day = month.date().replace(day=int(value))
                            if (day.weekday() + 1) % 7 != column - 2 or day in seen:
                                raise ValueError
                            days += 1
                            seen.add(day)
                            context[column] = day
                        except ValueError:
                            errors.append(cell.coordinate)
                            context.pop(column, None)
                    elif cell.data_type == "s" and isinstance(value, str) and value.strip():
                        if column in context:
                            texts += 1
                        else:
                            unassigned += 1
        return {"filename": parsed.filename, "digest": parsed.digest,
                "schedule_extent": schedule.calculate_dimension(), "calendar_extent": calendar.calculate_dimension(),
                "weekly_blocks": len(blocks), "schedule_dates": len(blocks) * len(DAY_COLUMNS),
                "direct_refs": direct, "mapped_refs": mapped, "invalid_refs": invalid,
                "crossed_refs": crossed, "indirect_refs": indirect,
                "months": len(month_rows), "days": days, "text_cells": texts,
                "unassigned_text_cells": unassigned, "grid_errors": errors,
                "exact_times": exact_times}
    finally:
        book.close()
