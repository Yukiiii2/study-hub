"""Literal ASSESSMENTS coverage extraction; no dates, results or helpers read."""
from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.parsers.curriculum_workbook import literal

SUBJECT_COLUMNS = {2: "FAR", 4: "AFAR", 6: "MAS", 8: "TAX", 10: "RFBT", 12: "AT", 14: "AP"}
SECTION_ROWS = (2, 11, 22, 34, 46)
MAPPING = {
    "sheet": "ASSESSMENTS", "subject_columns": SUBJECT_COLUMNS,
    "section_rows": SECTION_ROWS, "topic_match": "exact subject and whitespace-normalized title",
    "topic_aliases": {}, "whole_subject": "AP All topics in fifth section only",
    "source_key": "project-1:ASSESSMENTS:A<section-row>",
    "scheduled_at": None, "status": "planned", "results": "never imported",
}


@dataclass
class ParsedAssessments:
    filename: str
    digest: str
    sections: list[dict] = field(default_factory=list)
    skipped: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def parse_assessment_workbook(path: Path) -> ParsedAssessments:
    data = path.read_bytes()
    result = ParsedAssessments(path.name, sha256(data).hexdigest())
    book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    try:
        if "ASSESSMENTS" not in book.sheetnames:
            result.errors.append("Missing ASSESSMENTS worksheet")
            return result
        sheet = book["ASSESSMENTS"]
        for column, code in SUBJECT_COLUMNS.items():
            if literal(sheet.cell(1, column)) != code:
                result.errors.append(f"Unexpected subject header ASSESSMENTS!{sheet.cell(1, column).coordinate}")
        merges = {str(merge) for merge in sheet.merged_cells.ranges}
        expected_merges = {f"A{row}:N{row}" for row in SECTION_ROWS}
        if merges != expected_merges:
            result.errors.append("Unexpected ASSESSMENTS section merges; review source structure")
        for index, start in enumerate(SECTION_ROWS):
            cell = sheet.cell(start, 1)
            title = literal(cell)
            if not title or len(title) > 200:
                result.errors.append(f"Invalid literal section title ASSESSMENTS!A{start}")
                continue
            # Preserve literal title spelling; whitespace normalization is the only transformation.
            section = {"row": start, "title": title, "source_key": f"project-1:ASSESSMENTS:A{start}", "coverage": []}
            end = SECTION_ROWS[index + 1] if index + 1 < len(SECTION_ROWS) else sheet.max_row + 1
            for row in range(start + 1, end):
                for column, code in SUBJECT_COLUMNS.items():
                    source = sheet.cell(row, column)
                    if source.value is None or source.value == "":
                        continue
                    label = literal(source)
                    trace = {"section_row": start, "locator": f"ASSESSMENTS!{source.coordinate}", "subject": code}
                    if label is None:
                        # Do not expose raw formula/boolean/error values in manifests or reports.
                        result.skipped.append({**trace, "label": None, "reason": "nonliteral coverage cell excluded"})
                        continue
                    section["coverage"].append({**trace, "label": label,
                                                "whole_subject": start == 46 and code == "AP" and label == "All topics"})
            result.sections.append(section)
    finally:
        book.close()
    return result
