"""Safe Recall worksheet structure and exact existing-topic mapping; no history."""
from collections import Counter, defaultdict
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.parsers.curriculum_workbook import SUBJECT_SHEETS, literal, normalize

HEADERS = {"B1": "First Rep", "C2": "R1", "D2": "Rating", "E2": "R2", "F2": "Rating",
           "G2": "R3", "H2": "Rating", "I2": "R4", "J2": "Rating", "K2": "R5", "L2": "Rating",
           "M1": "Next Rep", "N1": "Final Rating", "O1": "Remarks"}
COLUMN_ROLES = {"B": "First Rep", "C": "R1", "D": "R1 rating", "E": "R2", "F": "R2 rating",
                "G": "R3", "H": "R3 rating", "I": "R4", "J": "R4 rating", "K": "R5",
                "L": "R5 rating", "M": "Next Rep", "N": "Final Rating", "O": "Remarks"}


def _kind(cell):
    if cell.value is None:
        return "blank"
    return {"f": "formula", "e": "error", "d": "date/time", "n": "numeric",
            "s": "text", "b": "boolean"}.get(cell.data_type, "other")


def inspect_recall(path: Path, topics: list[dict]) -> dict:
    source = path.read_bytes()
    result = {"filename": path.name, "digest": sha256(source).hexdigest(), "sheets": [],
              "mappings": [], "conflicts": [], "warnings": [], "existing_topics": len(topics),
              "phase4_context": [{"code": topic["code"], "title": normalize(topic["title"])}
                                 for topic in topics if topic["subject_code"] == "AP" and topic["code"] == "AP-04"]}
    by_title = defaultdict(list)
    for topic in topics:
        by_title[(topic["subject_code"], normalize(topic["title"]))].append(topic)
    book = load_workbook(BytesIO(source), data_only=False, read_only=True, keep_links=False)
    try:
        names = [name for name in book.sheetnames if name.endswith(" Recall")]
        for expected in SUBJECT_SHEETS:
            if expected + " Recall" not in names:
                result["warnings"].append(f"Missing expected worksheet: {expected} Recall.")
        for name in names:
            sheet = book[name]
            extent = sheet.calculate_dimension(force=True)
            subject = SUBJECT_SHEETS.get(name[:-7])
            invalid_headers = [coordinate for coordinate, label in HEADERS.items() if literal(sheet[coordinate]) != label]
            if invalid_headers:
                result["warnings"].append(f"{name}: unexpected structural headers at " + ", ".join(invalid_headers) + ".")
            if literal(sheet["A1"]) != "Topic":
                result["warnings"].append(f"{name}!A1: Topic header missing or different; column A is corroborated by the repetition layout.")
            if subject is None:
                result["warnings"].append(f"{name}: no approved subject-sheet mapping.")
            columns = {column: Counter() for column in COLUMN_ROLES}
            extra = Counter()
            types = Counter()
            seen_titles = set()
            inventory = {"sheet": name, "subject": subject, "extent": extent, "source_rows": 0,
                         "literal_titles": 0, "mapped": 0, "conflicts": 0, "duplicate_titles": 0,
                         "column_types": columns, "other_column_types": extra, "body_types": types}
            for cells in sheet.iter_rows(min_row=3):
                title_cell = cells[0]
                for index, cell in enumerate(cells[1:], start=2):
                    kind = _kind(cell)
                    types[kind] += 1
                    column = get_column_letter(index)
                    (columns[column] if column in columns else extra)[kind] += 1
                if title_cell.value is None:
                    continue
                inventory["source_rows"] += 1
                row = {"locator": f"{name}!{title_cell.coordinate}", "subject": subject}
                title = literal(title_cell)
                if not title:
                    reason = "formula/error or nonliteral topic title; never evaluated or copied"
                else:
                    inventory["literal_titles"] += 1
                    if title in seen_titles:
                        inventory["duplicate_titles"] += 1
                    seen_titles.add(title)
                    row["title"] = title
                    matches = by_title[(subject, title)] if subject else []
                    if invalid_headers or subject is None:
                        reason = "unverified sheet layout or subject mapping"
                    elif len(matches) != 1:
                        reason = "ambiguous existing normalized title" if matches else "no exact existing normalized title"
                    else:
                        result["mappings"].append({**row, "code": matches[0]["code"], "existing_title": normalize(matches[0]["title"])})
                        inventory["mapped"] += 1
                        continue
                result["conflicts"].append({**row, "reason": reason})
                inventory["conflicts"] += 1
            result["sheets"].append(inventory)
    finally:
        book.close()
    return result
