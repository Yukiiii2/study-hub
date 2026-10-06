"""Phase 5 extension of the reviewed curriculum parser; no completion data read."""
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.parsers.curriculum_workbook import ParsedWorkbook, SUBJECT_SHEETS, literal, parse_workbook

MAPPING = {
    "subject_sheets": SUBJECT_SHEETS,
    "topic_aliases": {"RFBT/RFBT-11A": "RFBT-11a"},
    "identity": "topic UUID + project-1:worksheet:row source locator",
    "display_order": "source row number",
    "duration_rules": {"h:mm": "MM:SS", "[h]:mm:ss": "elapsed MM:SS (stored final seconds must be zero)",
                       "h:mm:ss": "HH:MM:SS"},
    "confirmed_duration_corrections": {"FAR Vids!C294": "23:59", "FAR Vids!C295": "24:46", "TAX Vids!C64": "27:08"},
    "excluded_conflicts": ["FAR Vids!586", "FAR Vids!587"],
    "personal_progress": "not imported",
}


def duration_seconds(row: dict) -> int:
    """Only observed formats and explicitly confirmed cell corrections are accepted."""
    location = (row["sheet"], row["row"])
    raw, kind, fmt = row.get("duration_raw_xml_value"), row.get("duration_raw_xml_type"), row["duration_format"]
    corrections = {("FAR Vids", 294): ("23.59", 1439), ("FAR Vids", 295): ("24.46", 1486)}
    if location in corrections and fmt == "General" and kind == "n":
        expected, seconds = corrections[location]
        if raw == expected:
            return seconds
        raise ValueError("Confirmed source correction no longer matches")
    if location == ("TAX Vids", 64) and fmt == "General" and kind == "s" and row["duration_source"] == "27:08:":
        return 1628
    if kind != "n" or fmt not in MAPPING["duration_rules"]:
        raise ValueError("Unconfirmed duration format/value")
    try:
        stored = Decimal(raw) * 86400
        if not stored.is_finite() or stored < 0:
            raise ValueError("Invalid duration")
        whole = int(stored.to_integral_value(rounding=ROUND_HALF_UP))
        if abs(stored - whole) > Decimal("0.01"):
            raise ValueError("Fractional source seconds are unsupported")
    except (InvalidOperation, TypeError):
        raise ValueError("Invalid numeric duration") from None
    if fmt in ("h:mm", "[h]:mm:ss"):
        if whole % 60 or (fmt == "h:mm" and whole >= 86400):
            raise ValueError("Ambiguous MM:SS duration")
        seconds = whole // 60
        return seconds
    if whole >= 86400:
        raise ValueError("Time-of-day format would wrap hours")
    return whole


@dataclass
class ParsedVideos:
    curriculum: ParsedWorkbook
    videos: list[dict] = field(default_factory=list)
    candidates: Counter = field(default_factory=Counter)
    excluded: list[dict] = field(default_factory=list)
    skipped: list[dict] = field(default_factory=list)
    invalid_durations: list[dict] = field(default_factory=list)
    duplicates: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def parse_video_workbook(path: Path) -> ParsedVideos:
    curriculum = parse_workbook(path)
    result = ParsedVideos(curriculum)
    raw = {(r["sheet"], r["row"]): r for r in curriculum.deferred["videos"]}
    data = path.read_bytes()
    if sha256(data).hexdigest() != curriculum.digest:
        raise ValueError("Workbook changed during parsing")
    book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    try:
        for base, subject in SUBJECT_SHEETS.items():
            name = base + " Vids"
            if name not in book.sheetnames:
                result.errors.append(f"Missing video sheet {name}")
                continue
            ws = book[name]
            if [literal(ws.cell(1, c)) for c in (1, 2, 3)] != ["Topic #", "TOPICS", "Duration"]:
                result.errors.append(f"Unexpected video headers in {name}")
                continue
            for number in range(2, ws.max_row + 1):
                code_cell, title_cell = ws.cell(number, 1), ws.cell(number, 2)
                code, title = literal(code_cell), literal(title_cell)
                trace = {"sheet": name, "row": number, "subject": subject, "code": code, "title": title}
                if code_cell.data_type in ("f", "e"):
                    if title_cell.value is None and ws.cell(number, 3).value is None:
                        result.skipped.append({**trace, "reason": "formula-only summary with no lecture title/duration"})
                        continue
                    result.errors.append(f"Formula/error topic reference {name}!A{number}")
                    continue
                if not code:
                    if code_cell.value is not None:
                        result.errors.append(f"Invalid topic reference {name}!A{number}")
                    elif title_cell.value is not None:
                        result.skipped.append({**trace, "reason": "uncoded section/total label; not a lecture"})
                    continue
                result.candidates[subject] += 1
                if not title:
                    result.errors.append(f"Invalid title {name}!B{number}")
                    continue
                row = raw.get((name, number))
                if row is None or ws.cell(number, 3).data_type in ("f", "e"):
                    result.errors.append(f"Formula/error video duration {name}!C{number}")
                    continue
                if name == "FAR Vids" and number in (586, 587):
                    if code != "FAR-42" or title != "42-04 Exercise 2":
                        result.errors.append(f"Approved exclusion changed at {name}!{number}")
                    result.excluded.append({**trace, "reason": "unresolved source-data conflict; client/source clarification required", "duration_seconds": duration_seconds(row)})
                    continue
                try:
                    seconds = duration_seconds(row)
                except ValueError:
                    result.invalid_durations.append({**trace, "reason": "unconfirmed/malformed duration; skipped until source confirmed"})
                    continue
                result.videos.append({**trace, "duration_seconds": seconds, "display_order": number,
                                      "source_code": f"project-1:{name}:{number}", "source_url": None})
    finally:
        book.close()
    seen = set()
    for row in result.videos:
        key = (row["subject"], row["code"], row["title"].casefold())
        if key in seen:
            result.duplicates.append(row)
            result.errors.append(f"Duplicate source title on {row['sheet']}!{row['row']}; no automatic merge")
        seen.add(key)
    return result
