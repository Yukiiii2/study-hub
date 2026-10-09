"""Focused workbook/catalog/idempotency checks without database writes."""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from uuid import UUID

from openpyxl import Workbook

from app.parsers.assessment_workbook import SECTION_ROWS, SUBJECT_COLUMNS, parse_assessment_workbook
from app.services.assessment_import import make_plan, write_report


OWNER = UUID(int=100)


class AssessmentImportChecks(unittest.TestCase):
    def parse_source(self, extra=None):
        with TemporaryDirectory() as directory:
            book = Workbook()
            sheet = book.active
            sheet.title = "ASSESSMENTS"
            for column, subject in SUBJECT_COLUMNS.items():
                sheet.cell(1, column, subject)
            for index, row in enumerate(SECTION_ROWS, 1):
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=14)
                sheet.cell(row, 1, f"{index} MONTLY ASSESSMENT")
            sheet["B3"] = "Exact   topic"
            sheet["B4"] = "Unknown topic"
            sheet["B5"] = "=1+1"
            sheet["B6"] = True
            sheet["C3"] = "PRIVATE CHECKBOX HELPER"
            sheet["N47"] = "AP topic"
            sheet["N49"] = "All topics"
            if extra:
                extra(sheet)
            path = Path(directory) / "Project 1.xlsx"
            book.save(path)
            return parse_assessment_workbook(path)

    def catalog(self):
        subjects = [{"id": UUID(int=i + 1), "code": code, "is_active": True}
                    for i, code in enumerate(SUBJECT_COLUMNS.values())]
        ids = {s["code"]: s["id"] for s in subjects}
        topics = [{"id": UUID(int=30), "subject_id": ids["FAR"], "title": "Exact topic", "code": "FAR-1"},
                  {"id": UUID(int=31), "subject_id": ids["AP"], "title": "AP topic", "code": "AP-1"}]
        return subjects, topics

    def test_literal_parser_and_partial_coverage(self):
        parsed = self.parse_source()
        self.assertFalse(parsed.errors)
        self.assertEqual([s["row"] for s in parsed.sections], list(SECTION_ROWS))
        self.assertEqual(parsed.sections[0]["title"], "1 MONTLY ASSESSMENT")
        self.assertEqual(len(parsed.skipped), 2)
        subjects, topics = self.catalog()
        plan = make_plan(parsed, subjects, topics, [], OWNER)
        self.assertFalse(plan["errors"])
        self.assertEqual(len(plan["inserts"]), 5)
        first, fifth = plan["inserts"][0], plan["inserts"][4]
        self.assertEqual(first["topic_ids"], [str(topics[0]["id"])])
        self.assertIn("1 source coverage entries excluded", first["description"])
        self.assertIsNone(first["scheduled_at"])
        self.assertEqual(first["status"], "planned")
        self.assertEqual(fifth["subject_ids"], [str(subjects[-1]["id"])])
        self.assertEqual(fifth["topic_ids"], [])
        self.assertEqual(plan["unmapped"][0]["locator"], "ASSESSMENTS!B4")
        self.assertEqual(plan["unmapped"][0]["label"], "Unknown topic")
        self.assertEqual(len(plan["superseded"]), 1)

    def test_exact_matching_no_aliases_or_cross_subject(self):
        parsed = self.parse_source(lambda sheet: setattr(sheet["B3"], "value", "exact topic"))
        subjects, topics = self.catalog()
        topics.append({"id": UUID(int=32), "subject_id": subjects[-1]["id"], "title": "exact topic", "code": "AP-2"})
        plan = make_plan(parsed, subjects, topics, [], OWNER)
        self.assertEqual(plan["inserts"][0]["topic_ids"], [])
        self.assertEqual(len(plan["unmapped"]), 2)

    def test_idempotency_conflicts_and_owner_boundary(self):
        parsed = self.parse_source()
        subjects, topics = self.catalog()
        plan = make_plan(parsed, subjects, topics, [], OWNER)
        again = make_plan(parsed, subjects, topics, plan["inserts"], OWNER)
        self.assertEqual(again["inserts"], [])
        self.assertEqual(again["unchanged"], 5)
        self.assertFalse(again["errors"])
        for field, value in [("description", "edited"), ("scheduled_at", "2026-01-01"),
                             ("status", "completed"), ("topic_ids", []), ("title", "renamed")]:
            changed = deepcopy(plan["inserts"])
            changed[0][field] = value
            self.assertTrue(make_plan(parsed, subjects, topics, changed, OWNER)["errors"], field)
        unkeyed = deepcopy(plan["inserts"][:1])
        unkeyed[0]["source_key"] = None
        self.assertTrue(make_plan(parsed, subjects, topics, unkeyed, OWNER)["errors"])
        self.assertTrue(make_plan(parsed, subjects, topics, plan["inserts"] * 2, OWNER)["errors"])
        foreign = deepcopy(plan["inserts"])
        for row in foreign:
            row["user_id"] = UUID(int=101)
        other = make_plan(parsed, subjects, topics, foreign, OWNER)
        self.assertEqual(len(other["inserts"]), 5)
        self.assertFalse(other["errors"])

    def test_ambiguous_title_and_invalid_structure(self):
        subjects, topics = self.catalog()
        topics.append({**topics[0], "id": UUID(int=99)})
        plan = make_plan(self.parse_source(), subjects, topics, [], OWNER)
        self.assertEqual(plan["inserts"][0]["topic_ids"], [])
        self.assertIn("ambiguous", plan["unmapped"][0]["reason"])
        parsed = self.parse_source(lambda sheet: setattr(sheet["B1"], "value", "MS"))
        self.assertTrue(parsed.errors)
        subjects[0]["is_active"] = False
        self.assertTrue(make_plan(self.parse_source(), subjects, topics, [], OWNER)["errors"])

    def test_safe_reports_exclude_owner_and_helper_values(self):
        parsed = self.parse_source()
        subjects, topics = self.catalog()
        plan = make_plan(parsed, subjects, topics, [], OWNER)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "dry-run.md"
            write_report(parsed, plan, path)
            content = path.read_text(encoding="utf-8") + path.with_suffix(".plan.json").read_text(encoding="utf-8")
            self.assertNotIn(str(OWNER), content)
            self.assertNotIn("PRIVATE CHECKBOX HELPER", content)
            self.assertIn("ASSESSMENTS!B4", content)
            self.assertIn("Unknown topic", content)

    def test_real_source_exact_mapping_and_summary_exclusion(self):
        from app.parsers.curriculum_workbook import parse_workbook
        from app.services.curriculum_import import make_plan as curriculum_plan
        path = Path(__file__).resolve().parents[2] / "data/imports/Project 1.xlsx"
        if not path.exists():
            self.skipTest("Private workbook is local and not checked in")
        subjects, _ = self.catalog()
        topics = curriculum_plan(parse_workbook(path), subjects, [])["inserts"]
        parsed = parse_assessment_workbook(path)
        plan = make_plan(parsed, subjects, topics, [], OWNER)
        self.assertFalse(plan["errors"])
        self.assertEqual(len(parsed.skipped), 35)
        self.assertEqual(len(plan["unmapped"]), 74)
        self.assertEqual(len(plan["superseded"]), 2)
        self.assertEqual([s["topic_count"] for s in plan["sections"]], [13, 15, 12, 18, 13])
        self.assertEqual([s["excluded_count"] for s in plan["sections"]], [13, 15, 17, 13, 16])
        self.assertEqual([s["whole_subjects"] for s in plan["sections"]], [[], [], [], [], ["AP"]])
        self.assertTrue(all("MONTLY" in s["title"] for s in parsed.sections))


if __name__ == "__main__":
    unittest.main()
