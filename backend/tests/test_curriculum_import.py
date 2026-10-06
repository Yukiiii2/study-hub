"""Narrow migration checks. Fixtures never touch the configured database."""
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import UUID

from openpyxl import Workbook

from app.parsers.curriculum_workbook import parse_workbook, SUBJECT_SHEETS
from app.services.curriculum_import import make_plan, validate_integrity

SUBJECTS = [{"id": UUID(int=i + 1), "code": code, "is_active": True}
            for i, code in enumerate(SUBJECT_SHEETS.values())]


class ImportChecks(unittest.TestCase):
    def workbook(self, path):
        book = Workbook()
        book.remove(book.active)
        for sheet, code in SUBJECT_SHEETS.items():
            ws = book.create_sheet(sheet)
            ws["B2"] = "HO#"
            ws["C2"] = "Sub-topics"
            ws["B4"] = code + "-01"
            ws["C4"] = "  A real\nsource title  "
            mirror = book.create_sheet(sheet + " Code")
            mirror["A2"], mirror["B2"] = "CODE", "TOPIC"
            mirror["A3"], mirror["B3"] = code + "-01", "A real source title"
        book.save(path)

    def test_formula_titles_block_import_and_blank_titles_are_excluded(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            self.workbook(path)
            from openpyxl import load_workbook
            book = load_workbook(path)
            book["FAR"]["B5"], book["FAR"]["C5"] = "FAR-19", None
            book["AT"]["C4"] = '=CONCAT("invented")'
            book.save(path)
            parsed = parse_workbook(path)
        self.assertEqual(len(parsed.topics), 6)
        self.assertTrue(any("formula" in error.lower() for error in parsed.errors))
        self.assertTrue(any(item["code"] == "FAR-19" for item in parsed.ambiguous))

    def test_idempotency_and_conflicts(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            self.workbook(path)
            parsed = parse_workbook(path)
        first = make_plan(parsed, SUBJECTS, [])
        self.assertFalse(first["errors"])
        self.assertEqual(len(first["inserts"]), 7)
        second = make_plan(parsed, SUBJECTS, first["inserts"])
        self.assertFalse(second["errors"])
        self.assertEqual(len(second["inserts"]), 0)
        self.assertEqual(second["unchanged"], 7)
        changed = [dict(row) for row in first["inserts"]]
        changed[0]["title"] = "Different existing title"
        self.assertTrue(make_plan(parsed, SUBJECTS, changed)["errors"])
        duplicated = first["inserts"] + [first["inserts"][0]]
        self.assertTrue(make_plan(parsed, SUBJECTS, duplicated)["errors"])

    def test_duplicate_source_codes_and_missing_subject_block_import(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            self.workbook(path)
            from openpyxl import load_workbook
            book = load_workbook(path)
            book["MS"]["B5"], book["MS"]["C5"] = "MAS-01", "Another title"
            book.save(path)
            parsed = parse_workbook(path)
        self.assertTrue(parsed.errors)
        self.assertTrue(make_plan(parsed, SUBJECTS[1:], [])["errors"])

    def test_invalid_parent_integrity_blocks_import(self):
        base = {"id": "one", "subject_id": "subject", "parent_topic_id": "one", "code": "CODE"}
        self.assertTrue(validate_integrity([base]))
        self.assertTrue(validate_integrity([{**base, "parent_topic_id": "missing"}]))
        self.assertTrue(validate_integrity([{**base, "parent_topic_id": "two"},
                                            {**base, "id": "two", "subject_id": "other", "parent_topic_id": None, "code": "OTHER"}]))


if __name__ == "__main__":
    unittest.main()
